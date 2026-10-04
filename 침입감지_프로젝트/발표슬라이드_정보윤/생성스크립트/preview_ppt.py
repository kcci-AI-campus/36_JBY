# -*- coding: utf-8 -*-
# pptx 를 matplotlib 로 근사 렌더링해 PNG 미리보기를 만든다 (PowerPoint 없이 배치·넘침 확인용).
import sys, os, io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from matplotlib.patches import FancyBboxPatch, Rectangle
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
fm.fontManager.addfont(r"C:\Windows\Fonts\malgun.ttf"); fm.fontManager.addfont(r"C:\Windows\Fonts\malgunbd.ttf")
plt.rcParams["font.family"] = "Malgun Gothic"; plt.rcParams["axes.unicode_minus"] = False
SRC = sys.argv[1]
OUTDIR = sys.argv[2]
os.makedirs(OUTDIR, exist_ok=True)
prs = Presentation(SRC)
SW, SH = prs.slide_width / 914400, prs.slide_height / 914400
E = 914400.0


def cw(ch, size):  # 글자 폭(인치) 근사
    o = ord(ch)
    if o < 128: return size * (0.28 if ch in " .,:;'|" else 0.56) / 72
    if 0x2000 <= o < 0x2100: return size * 0.5 / 72
    return size * 1.0 / 72


def wrap(txt, size, width):
    lines = []
    for raw in txt.split("\n"):
        cur, w = "", 0.0
        for ch in raw:
            cwid = cw(ch, size)
            if w + cwid > width and cur:
                lines.append(cur); cur, w = "", 0.0
            cur += ch; w += cwid
        lines.append(cur)
    return lines


def color_of(font, default="#475569"):
    try:
        if font.color and font.color.type is not None and font.color.rgb is not None:
            return "#" + str(font.color.rgb)
    except Exception:
        pass
    return default


def draw_tf(ax, tf, l, t, w, h, default_size=11, pad=0.03):
    """text frame 을 그린다. 문단마다 첫 run 의 크기·색·굵기를 쓴다(우리 슬라이드는 문단 안 run 이 같은 크기)."""
    try:
        ml = (tf.margin_left or 0) / E; mr = (tf.margin_right or 0) / E; mt = (tf.margin_top or 0) / E
    except Exception:
        ml = mr = mt = pad
    inner_w = max(0.1, w - ml - mr)
    items = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs]
        if not runs:
            items.append([("", default_size, False, "#000000", 1.0, p.alignment)]); continue
        size = runs[0].font.size.pt if runs[0].font.size else default_size
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        # run 별로 폭을 재서 한 문단을 이어 붙인다 (색은 run 별)
        segs = [(r.text, r.font.size.pt if r.font.size else size, bool(r.font.bold), color_of(r.font)) for r in runs]
        items.append((segs, size, ls, p.alignment))
    # 총 높이 계산
    total = 0.0; laid = []
    for it in items:
        if isinstance(it, list): laid.append(("", 1, [], PP_ALIGN.LEFT, default_size)); total += default_size * 1.0 / 72; continue
        segs, size, ls, al = it
        text = "".join(s[0] for s in segs)
        lines = wrap(text, size, inner_w)
        lh = size * ls * 1.18 / 72
        laid.append((lines, lh, segs, al, size)); total += lh * len(lines)
    anchor = tf.vertical_anchor
    y = t + mt
    if anchor == MSO_ANCHOR.MIDDLE: y = t + (h - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM: y = t + h - total - mt
    for lines, lh, segs, al, size in laid:
        if not segs: y += lh; continue
        # 색: 문단 안 run 이 여러 색이면 첫 줄 앞부분만 첫 run 색 (근사)
        col = segs[0][3]; bold = segs[0][2]
        for k, ln in enumerate(lines):
            x = l + ml
            lw = sum(cw(c, size) for c in ln)
            if al == PP_ALIGN.CENTER: x = l + ml + (inner_w - lw) / 2
            elif al == PP_ALIGN.RIGHT: x = l + w - mr - lw
            # 첫 줄에 run 이 둘 이상이면 run 별 색으로 나눠 그린다
            if k == 0 and len(segs) > 1 and "".join(s[0] for s in segs) == ln:
                xx = x
                for st, ss, sb, sc in segs:
                    ax.text(xx, y + lh * 0.55, st, fontsize=ss * 0.98, color=sc, fontweight="bold" if sb else "normal", va="center", ha="left")
                    xx += sum(cw(c, ss) for c in st)
            else:
                ax.text(x, y + lh * 0.55, ln, fontsize=size * 0.98, color=col, fontweight="bold" if bold else "normal", va="center", ha="left")
            y += lh
    overflow = (y - t) - h
    return overflow


report = []
for idx, slide in enumerate(prs.slides, 1):
    fig = plt.figure(figsize=(SW, SH), dpi=160)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, SW); ax.set_ylim(SH, 0); ax.axis("off")
    ax.add_patch(Rectangle((0, 0), SW, SH, fc="white", ec="none"))
    for sh in slide.shapes:
        l, t, w, h = sh.left / E, sh.top / E, sh.width / E, sh.height / E
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = Image.open(io.BytesIO(sh.image.blob))
            ax.imshow(img, extent=(l, l + w, t + h, t), aspect="auto", zorder=2)
            continue
        if sh.has_table:
            tbl = sh.table
            ys = t
            for i, row in enumerate(tbl.rows):
                xs = l; rh = row.height / E
                for j, col in enumerate(tbl.columns):
                    cwid = col.width / E
                    cell = tbl.cell(i, j)
                    fc = "#FFFFFF"
                    try:
                        if cell.fill.type == 1: fc = "#" + str(cell.fill.fore_color.rgb)
                    except Exception: pass
                    ax.add_patch(Rectangle((xs, ys), cwid, rh, fc=fc, ec="#E2E8F0", lw=0.4, zorder=1))
                    ov = draw_tf(ax, cell.text_frame, xs, ys, cwid, rh, default_size=8)
                    if ov > 0.02: report.append("slide %d table cell (%d,%d) overflow %.2f in: %s" % (idx, i, j, ov, cell.text_frame.text[:30]))
                    xs += cwid
                ys += rh
            continue
        if sh.shape_type in (MSO_SHAPE_TYPE.LINE,) or sh.__class__.__name__ == "Connector":
            col = "#94A3B8"
            try: col = "#" + str(sh.line.color.rgb)
            except Exception: pass
            ax.annotate("", xy=(sh.end_x / E, sh.end_y / E), xytext=(sh.begin_x / E, sh.begin_y / E),
                        arrowprops=dict(arrowstyle="-|>", color=col, lw=1.2, shrinkA=0, shrinkB=0), zorder=3)
            continue
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            ax.add_patch(Rectangle((l, t), w, h, fc="none", ec="#CBD5E1", lw=0.5, ls="--", zorder=2))
            try:
                ct = str(ch.chart_type)
                if "XY" in ct:
                    for ser in ch.plots[0].series:
                        xs = list(ser.iter_values()) if False else None
                    import re as _re
                    xml = ch._chartSpace.xml
                    for sx in ch._chartSpace.xpath(".//c:ser"):
                        xv = [float(v) for v in sx.xpath("./c:xVal//c:v/text()")]
                        yv = [float(v) for v in sx.xpath("./c:yVal//c:v/text()")]
                        isline = not sx.xpath("./c:marker/c:symbol[@val='square']")
                        X = [l + 0.3 + (w - 0.4) * x / 60.0 for x in xv]; Y = [t + h - 0.15 - (h - 0.25) * y / 100.0 for y in yv]
                        if isline: ax.plot(X, Y, color="#2563EB", lw=0.5, zorder=3)
                        else: ax.scatter(X, Y, s=0.6, color="#F97316", zorder=4)
                else:
                    sers = ch._chartSpace.xpath(".//c:ser")
                    cats = len(sers[0].xpath("./c:val//c:pt"))
                    cols = ["#94A3B8", "#2563EB", "#F97316"]
                    for j, sx in enumerate(sers):
                        vals = [float(v) for v in sx.xpath("./c:val//c:v/text()")]
                        for c_, v in enumerate(vals):
                            gw = (w - 0.3) / cats; bw = gw * 0.7 / len(sers)
                            x0 = l + 0.15 + c_ * gw + gw * 0.15 + j * bw
                            bh = (h - 0.55) * v / 8.5
                            ax.add_patch(Rectangle((x0, t + h - 0.35 - bh), bw * 0.95, bh, fc=cols[j % 3], zorder=3))
                            ax.text(x0 + bw / 2, t + h - 0.37 - bh, "%.1f" % v, fontsize=5, ha="center", va="bottom", zorder=4)
            except Exception as ex:
                ax.text(l, t, "chart " + str(ex)[:40], fontsize=5)
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fc = None
            try:
                if sh.fill.type == 1: fc = "#" + str(sh.fill.fore_color.rgb)
            except Exception: pass
            ec = "none"
            try:
                if sh.line.fill.type == 1: ec = "#" + str(sh.line.color.rgb)
            except Exception: pass
            st = str(sh.auto_shape_type)
            if "DIAMOND" in st:
                from matplotlib.patches import Polygon as _Poly
                ax.add_patch(_Poly([[l + w / 2, t], [l + w, t + h / 2], [l + w / 2, t + h], [l, t + h / 2]], fc=fc or "white", ec=ec, lw=1.2, zorder=2))
                fc = None
            elif "OVAL" in st:
                from matplotlib.patches import Ellipse as _Ell
                ax.add_patch(_Ell((l + w / 2, t + h / 2), w, h, fc=fc or "white", ec=ec, lw=0.8, zorder=4))
                fc = None
                if sh.has_text_frame:
                    ax.text(l + w / 2, t + h / 2, sh.text_frame.text, fontsize=6.5, color="white", ha="center", va="center", fontweight="bold", zorder=5)
                continue
            if fc:
                rs = 0.06
                try: rs = sh.adjustments[0] * min(w, h) if len(sh.adjustments) else 0.0
                except Exception: rs = 0.0
                if rs > 0.005:
                    ax.add_patch(FancyBboxPatch((l + rs, t + rs), w - 2 * rs, h - 2 * rs, boxstyle="round,pad=%f" % rs, fc=fc, ec=ec, lw=1.2, zorder=1 if ec == "none" else 2))
                else:
                    ax.add_patch(Rectangle((l, t), w, h, fc=fc, ec=ec, zorder=1))
        if sh.has_text_frame and sh.text_frame.text.strip():
            ov = draw_tf(ax, sh.text_frame, l, t, w, h)
            if ov > 0.03: report.append("slide %d text overflow %.2f in: %s" % (idx, ov, sh.text_frame.text[:40].replace("\n", " ")))
    p = os.path.join(OUTDIR, "slide%d.png" % idx)
    fig.savefig(p, dpi=160); plt.close(fig)
    print("rendered", p)
print("\n".join(report) if report else "no overflow detected (근사)")
