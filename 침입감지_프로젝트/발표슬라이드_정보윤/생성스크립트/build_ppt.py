# -*- coding: utf-8 -*-
# 정보윤 파트 발표 슬라이드 생성: 팀장 디자인(파이썬 프로젝트 발표자료.pptx)의 마스터·테마를 그대로 쓰고 슬라이드만 새로 만든다.
import sys, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
BASE = r"C:\Users\LG\Developer\침입감지_프로젝트"
FIG = os.path.join(BASE, "발표슬라이드_정보윤")
SRC = os.path.join(BASE, "파이썬 프로젝트 발표자료.pptx")
OUT = os.path.join(FIG, "발표_정보윤파트_0927_v5.pptx")
FONT = "Noto Sans KR"
C = dict(blue="2563EB", navy="0F172A", gray="475569", muted="64748B", panel="F1F5F9", white="FFFFFF",
         red="DC2626", green="16A34A", amber="D97706", slate="334155", band="F8FAFC")


def rgb(h): return RGBColor.from_string(h)


prs = Presentation(SRC)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    rId = sldId.rId
    sldIdLst.remove(sldId)
    prs.part.drop_rel(rId)
BLANK = [l for l in prs.slide_layouts if l.name == "BLANK"][0]


def set_ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def panel(s, l, t, w, h, fill="F1F5F9", adj=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.adjustments[0] = adj
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def bar(s, l, t, w, h, fill="2563EB"):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def text(s, l, t, w, h, paras, size=11, bold=False, color="475569", align="l", anchor="t", spacing=1.1, margin=0.03):
    """paras: 문자열 | 문단 리스트. 문단 = 문자열 | run 리스트. run = (text[, size[, bold[, color]]])"""
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            if isinstance(r, str): r = (r,)
            rt = r[0]; rs = r[1] if len(r) > 1 else size; rb = r[2] if len(r) > 2 else bold; rc = r[3] if len(r) > 3 else color
            run = para.add_run(); run.text = rt.replace("−", "-")
            f = run.font; f.name = FONT; f.size = Pt(rs); f.bold = rb; f.color.rgb = rgb(rc)
            set_ea(run)
    return tb


def pic_fit(s, path, l, t, w, h, align="c"):
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih); pw, ph = iw * sc, ih * sc
    dl = l + (w - pw) / 2 if align == "c" else l
    dt = t + (h - ph) / 2
    return s.shapes.add_picture(path, Inches(dl), Inches(dt), Inches(pw), Inches(ph))


import native
native.init(globals())

NOGRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def table(s, rows, l, t, col_w, row_h=0.3, size=8.5, hdr_size=8.5, hdr_fill="0F172A", hdr_color="FFFFFF",
          band=("FFFFFF", "F1F5F9"), bold_first_col=True, aligns=None, heights=None):
    nr, nc = len(rows), len(rows[0])
    hs = heights or [row_h] * nr
    gs = s.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(sum(col_w)), Inches(sum(hs)))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None: style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = NOGRID
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(hs[i])
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.05); cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(hdr_fill if i == 0 else band[(i - 1) % 2])
            val = rows[i][j]
            if isinstance(val, str): txt, col, bd = val, None, None
            else: txt = val[0]; col = val[1] if len(val) > 1 else None; bd = val[2] if len(val) > 2 else None
            tf = cell.text_frame; tf.word_wrap = True
            lines = txt.split("\n")
            for k, ln in enumerate(lines):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER if (i == 0 or (aligns and aligns[j] == "c")) else PP_ALIGN.LEFT
                run = p.add_run(); run.text = ln.replace("−", "-")
                f = run.font; f.name = FONT; f.size = Pt(hdr_size if i == 0 else size)
                f.bold = True if i == 0 else (bd if bd is not None else (bold_first_col and j == 0))
                f.color.rgb = rgb(hdr_color if i == 0 else (col or C["gray"]))
                set_ea(run)
    return gs


def header(s, label, title, conclusion=None):
    text(s, 0.23, 0.2, 7.5, 0.18, [[(label, 9.77, True, C["blue"])]], spacing=1.0)
    text(s, 0.23, 0.4, 9.5, 0.4, [[(title, 21.98, True, C["navy"])]], spacing=1.0)
    if conclusion:
        bar(s, 0.23, 0.86, 0.06, 0.3)
        text(s, 0.34, 0.83, 9.4, 0.36, [[("결론  ", 11.5, True, C["blue"]), (conclusion, 11.5, True, C["navy"])]],
             anchor="m", spacing=1.0)


def note(s, txt, t=5.22, h=0.32):
    text(s, 0.23, t, 9.54, h, [[(txt, 7.8, False, C["muted"])]], spacing=1.0)


def item(s, l, t, w, h, num, title, sub, fill="F1F5F9"):
    panel(s, l, t, w, h, fill, adj=0.14)
    text(s, l + 0.16, t + 0.1, 0.5, 0.26, [[(num, 13.43, True, C["blue"])]], spacing=1.0)
    text(s, l + 0.6, t + 0.09, w - 0.7, 0.24, [[(title, 10.99, True, C["navy"])]], spacing=1.0)
    text(s, l + 0.6, t + 0.34, w - 0.7, h - 0.38, [[(sub, 9.16, False, C["muted"])]], spacing=1.05)


def hero(s, l, t, w, h, number, caption, ncolor="2563EB", fill="F1F5F9", nsize=15):
    panel(s, l, t, w, h, fill, adj=0.12)
    text(s, l + 0.06, t + 0.06, w - 0.12, 0.4, [[(number, nsize, True, ncolor)]], align="c", anchor="m", spacing=1.0)
    text(s, l + 0.06, t + 0.46, w - 0.12, h - 0.48, [[(caption, 8, False, C["gray"])]], align="c", spacing=1.0)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


LABEL33 = "CORE TECHNOLOGY ③  알림 파이프라인 · 저전력 게이팅  —  정보윤"
SCRIPTS = {}

# ═══════════════════ 슬라이드 1: 알림 파이프라인 ═══════════════════
s = prs.slides.add_slide(BLANK)
header(s, LABEL33, "알림 파이프라인 — 판정부터 휴대폰까지 보드 한 대에서",
       "위험 판정 → 보드 LED → 휴대폰 알림음·사진 → 탭하면 실시간 화면. 평상시 영상은 밖으로 나가지 않는다")
panel(s, 0.23, 1.25, 9.54, 3.05, C["white"], adj=0.03)
native.pipeline(s)
hero(s, 0.23, 4.34, 2.0, 0.84, "0.5초 · 60 % · 3장", "체류 판정: 창 길이 · 위험 비율 · 최소 표본", nsize=12.5)
hero(s, 2.3, 4.34, 1.95, 0.84, "30초", "같은 경보의 휴대폰 알림 최소 간격 (보드 LED는 매번)")
hero(s, 4.32, 4.34, 1.95, 0.84, "앞 3초 + 뒤 2초", "경보 클립 길이 · 보드에 최근 50개만 보관", nsize=13.5)
text(s, 6.4, 4.3, 3.37, 0.2, [[("④ 알림을 탭하면 열리는 실시간 화면 (터널 주소 + 접속 암호)", 8.3, True, C["navy"])]], spacing=1.0)
pic_fit(s, os.path.join(FIG, "웹대시보드_상단.png"), 6.4, 4.5, 3.37, 0.68, align="l")
note(s, "체류 판정·반경 계수·알림 간격은 실행 옵션으로 조정 가능 · 경보 → 휴대폰 도착 지연은 월요일(9/28) 실측 예정 "
        "(구조상 체류 0.5 + 추론 0.02~0.11 + 전송 1~2초 ≈ 1.5~2.6초)")
SCRIPTS[1] = ("3-3 알림 파이프라인", """3-3 핵심 기술, 알림 파이프라인과 저전력 게이팅은 제가 맡았습니다. 먼저 알림 파이프라인입니다.
결론부터 말씀드리면, 카메라 영상이 들어와서 위험을 판정하고 휴대폰까지 알리는 과정이 전부 라즈베리파이 한 대 안에서 끝나고, 평상시 영상은 밖으로 나가지 않습니다. 경보 순간의 사진과 5초 클립만 나갑니다.

그림의 ①번, USB 웹캠의 640×480 영상을 YOLO11n이 아기·어른·칼·콘센트 네 가지로 검출합니다.
②번 거리 판정은 위험물 박스 폭의 1.5배를 위험 반경으로 잡고, 아기 박스가 그 반경에 닿으면 '위험'으로 봅니다. 폭에 비례시킨 이유는 멀리 있는 칼은 화면에서 작게 보이니 반경도 같이 작아지게 하려는 것입니다.
③번 체류 판정이 오경보를 막는 핵심입니다. 한 장에서 위험이 잡혔다고 바로 울리지 않고, 최근 0.5초 프레임의 60 % 이상이 위험일 때만 경보를 냅니다. 오검출은 보통 한두 장 깜빡이고 사라지기 때문에 이 규칙으로 걸러집니다. 최소 3장을 보는 이유는 2장으로는 다수결이 안 되기 때문입니다.
④번 경보가 나면 보드의 LED가 즉시 빨간색으로 바뀌고, ntfy라는 무료 푸시 서비스로 휴대폰에 알림음과 함께 사진 1장이 가고, 이어서 경보 앞 3초·뒤 2초 클립이 갑니다. 알림을 탭하면 오른쪽 아래 화면처럼 실시간 웹 페이지가 열립니다. 외부에서는 Cloudflare 터널 주소와 접속 암호가 있어야만 열립니다.

같은 경보가 계속되면 휴대폰 알림은 최소 30초 간격으로만 보내서 알림 피로를 줄였고, 보드 LED는 매번 켜집니다.
경보에서 휴대폰 도착까지 걸리는 시간은 아직 실측 전입니다. 구조상 체류 0.5초에 전송 1~2초를 더한 값이라 월요일에 스톱워치로 재서 숫자를 넣겠습니다.

[발표자 참고 — 말하지 않아도 되는 메모] 소리는 휴대폰 알림음이다. 보드 경보음 코드(alarm.wav를 aplay로 재생)는 있지만 라즈베리파이 5에는 오디오 단자가 없어 USB·블루투스 스피커와 alarm.wav 파일이 있어야 난다. 지금 보드들에는 둘 다 없다. "현장 경보음은?"이 나오면 "코드는 있고 스피커만 달면 된다, 향후 과제"라고 답할 것. 파이프라인 코드는 우리 판(detect_baby.py, notify.py, clip.py, web.py) 기준. 터널·암호·알림 동작은 9/27 집(지인) 보드에서 클로드코드로 확인한 것. 질문이 나오면 "0.5초·60 %·3장은 절충값이며 옵션으로 바꿀 수 있고 월요일 0.3/0.5/1.0초를 비교해 확정하겠다"고 먼저 말할 것.""")

# ═══════════════════ 슬라이드 2: 움직임 게이팅 ═══════════════════
s = prs.slides.add_slide(BLANK)
header(s, LABEL33, "움직임 게이팅 — 화면이 안 변하면 추론을 건너뛴다",
       "같은 장면·같은 모델에서 7.12 → 2.52 W (−65 %), 추론 749 → 32회/60초. 위험 중에는 절대 건너뛰지 않는다")
panel(s, 0.23, 1.25, 3.35, 3.15, C["white"], adj=0.03)
native.gating(s)
table(s, [["항목", "게이팅 끔", "게이팅 켬"],
          ["평균 전력", "7.12 W", ("2.52 W  (−65 %)", C["blue"], True)],
          ["CPU 사용률", "92 %", ("5 %", C["blue"], True)],
          ["추론 횟수 / 60초", "749회", ("32회  (93 % 건너뜀)", C["blue"], True)]],
      0.23, 4.47, [1.15, 0.8, 1.4], row_h=0.2, size=8, hdr_size=8, aligns=["l", "c", "c"])
native.timeline(s, BASE)
panel(s, 3.68, 4.47, 6.09, 0.7, C["panel"], adj=0.12)
text(s, 3.8, 4.5, 5.9, 0.22, [[("다른 보드·다른 장면에서도 재현 — 교육장 보드 9/23, 인형·칼을 놓은 정지 장면", 8.8, True, C["navy"])]], spacing=1.0)
text(s, 3.8, 4.75, 1.95, 0.38, [[("6.03 → 2.26 W", 13, True, C["blue"])], [("전력 −62 %", 7.8, False, C["gray"])]], align="c", spacing=1.0)
text(s, 5.75, 4.75, 1.95, 0.38, [[("91 → 7~16 %", 13, True, C["blue"])], [("CPU 사용률", 7.8, False, C["gray"])]], align="c", spacing=1.0)
text(s, 7.7, 4.75, 1.95, 0.38, [[("57.9 → 45.8 ℃", 13, True, C["blue"])], [("보드 온도", 7.8, False, C["gray"])]], align="c", spacing=1.0)
note(s, "측정: 라즈베리파이 5 PMIC(USB 웹캠 전력 제외) · yolo11n 512 FP32 · 실제 카메라, 빈 방 60초 · 2026-09-26 지인 보드 | 안전장치: 위험 중 건너뛰기 금지, "
        "연속 15장이면 강제 추론, 쉴 때 8 FPS(최대 공백 0.13초)")
SCRIPTS[2] = ("3-3 움직임 게이팅", """다음은 저전력의 첫 번째 방법, 움직임 게이팅입니다.
결론은, 같은 장면·같은 모델에서 평균 전력이 7.12 W에서 2.52 W로 65 % 줄었고, 60초 동안의 추론 횟수가 749회에서 32회로 줄었다는 것입니다.

원리는 왼쪽 흐름도입니다. A, 새 프레임이 오면 직전 프레임과의 차이를 160×120 흑백으로 줄여서 계산합니다. B, 변한 픽셀이 0.6 % 미만이고 지금 위험 중이 아니면 추론을 건너뛰고 직전 결과를 그대로 씁니다. 화면이 안 변했으면 답도 안 변했을 테니까요. 원래 기획에서는 초음파 센서로 '가까이 왔을 때만 추론'하려 했는데, 센서 없이 소프트웨어로 같은 효과를 낸 것입니다.

그냥 건너뛰기만 하면 위험하니까 안전장치 세 개를 넣었습니다(C). 첫째, 위험 판정 중에는 절대 건너뛰지 않습니다. 둘째, 연속 15장을 건너뛰면 강제로 한 번 추론합니다. 셋째, 쉴 때도 초당 8장은 확인하므로 최대 공백은 0.13초입니다. 체류 판정이 0.5초를 보기 때문에 이 공백은 판정에 영향이 없습니다.

오른쪽 그래프가 실측입니다. 위가 게이팅 끔인데, 주황 점이 추론이 실행된 순간이고 빈틈없이 돌아서 CPU가 92 %입니다. 아래가 켬인데, 추론이 약 2초에 한 번(15장 강제 추론)만 뛰고 CPU는 5 %입니다. 라즈베리파이 5의 전원 칩 PMIC로 잰 값이라 USB 웹캠 전력은 빠져 있고, 9월 26일 실제 카메라로 빈 방 60초를 잰 값입니다.
교육장 보드에서 9월 23일 인형과 칼을 놓은 정지 장면으로도 6.03에서 2.26 W, 62 % 감소로 재현됐고, CPU 91 %가 7~16 %, 보드 온도도 58도에서 46도로 내려갔습니다.

잃는 것은 쉴 때 최대 0.13초 지연 하나뿐이라 기본값으로 켜 두었습니다. 다만 아기가 계속 움직이는 장면에서는 매 프레임 추론해야 해서 게이팅 효과가 작습니다. 그래서 다음 슬라이드의 상태 설계가 필요합니다.

[발표자 참고] 9/26 측정은 클로드코드가 지인 보드에서 실행한 것, 9/23 교육장 측정은 본인이 실행. "게이팅 켜면 사고 순간을 놓치지 않나?"가 나오면 안전장치 세 개(위험 중 금지·15장 강제·8 FPS)를 그대로 답하면 됨.""")

# ═══════════════════ 슬라이드 3 (선택): 4상태 ═══════════════════
s = prs.slides.add_slide(BLANK)
header(s, LABEL33 + "   (선택 슬라이드)", "상황에 따라 추론 속도를 바꾼다 — 두 모드에서 4상태로",
       "위험이 멀수록 천천히·싸게, 가까울수록 빠르게·정확하게 → 아기 움직일 때 4.3 → 1.8 W, 경보 구간 7.2 → 2.6 W")
panel(s, 0.23, 1.25, 6.45, 2.25, C["white"], adj=0.03)
native.states(s)
panel(s, 6.8, 1.25, 2.97, 2.25, C["panel"], adj=0.08)
text(s, 6.92, 1.32, 2.75, 0.22, [[("상태 수는 어떻게 정했나", 10, True, C["navy"])]], spacing=1.0)
text(s, 6.92, 1.55, 2.75, 0.5, [[("안전 = 상태별 최소 FPS 제약 (경보 ≥ 6, 경계 ≥ 12, 감시 5). 상태 수 = 하나 더 나눠 얻는 전력이 측정 잡음 0.1 W 보다 큰가", 8.2, False, C["gray"])]], spacing=1.05)
table(s, [["나눔", "전력 이득", "판단"],
          ["3 → 4상태\n감시·경계 분리", "-1.1 W\n(칼 멀리 있을 때)", ("채택", C["green"], True)],
          ["4 → 5상태\n원거리 2 FPS 추가", "-0.08 W", ("잡음 이하\n→ 안 나눔", C["red"], True)]],
      6.92, 2.1, [1.1, 0.9, 0.75], size=7.6, hdr_size=7.8, aligns=["l", "c", "c"], heights=[0.22, 0.5, 0.5])
table(s, [["장면 (320 입력, 45~60초)", "두 모드\n(팀원A 원본)", "4상태\n(FP32, 2400 MHz)", "4상태 + INT8\n+ 클럭 1500"],
          ["아기 가만히, 위험물 없음", "4.21 W", ("1.61 W", C["blue"], True), "미측정"],
          ["아기 움직임, 위험물 없음", "4.29 W", ("2.53 W", C["blue"], True), ("1.78 W", C["blue"], True)],
          ["아기 + 칼 멀리 (반경×2.8)", "미측정", ("1.63 W", C["blue"], True), "미측정"],
          ["아기 + 칼 반경 안 (경보 구간)", "7.17 W", "6.66 W", ("2.58 W", C["blue"], True)],
          ["빈 방, 실제 카메라 (대기)", "1.82 W", "1.80 W", ("1.74 W", C["blue"], True)]],
      0.23, 3.6, [2.2, 1.2, 1.35, 1.35], size=8, hdr_size=7.8, aligns=["l", "c", "c", "c"], heights=[0.36, 0.24, 0.24, 0.24, 0.24, 0.24])
panel(s, 6.5, 3.6, 3.27, 1.56, C["white"], adj=0.03)
native.bars(s)
note(s, "측정: 라즈베리파이 5 PMIC(웹캠 제외), 지인 보드, 2026-09-26~27 · 아기 장면은 9/23 클립으로 만든 영상 입력(근사치) → 월요일 실물 재측정 · "
        "경보 구간: 4상태 두 열은 30 FPS 상한, 두 모드는 상한 없음 · INT8 = 최신 120에폭 가중치, FP32와 같은 물체 검출 확인·신뢰도 0.2~0.3 낮음 · 아무것도 안 돌릴 때 1.55 W")
SCRIPTS[3] = ("3-3 (선택) 4상태 저전력", """게이팅은 '화면이 안 변할 때'만 효과가 있습니다. 아기가 움직이면 매 프레임 추론해야 해서 전력이 다시 올라갑니다. 그래서 두 번째 방법으로, 상황에 따라 추론 속도를 바꾸는 상태 설계를 했습니다.

원리는 한 줄입니다. 전력은 유휴 전력에, 추론 1회 에너지 곱하기 초당 추론 횟수를 더한 것입니다. 그러니 줄일 곳은 '몇 번 하느냐'와 '한 번을 얼마나 싸게 하느냐' 둘뿐입니다.
팀원A님의 두 모드(대기·고성능)를 출발점으로, 위험이 멀수록 천천히·싸게, 가까울수록 빠르게·정확하게 보도록 네 상태로 나눴습니다.
① 대기: 아기가 없으면 5초에 한 장. ② 감시: 아기는 있지만 위험물이 멀면 5 FPS에 게이팅, CPU 클럭도 1500 MHz로 낮춥니다. ③ 경계: 아기와 위험물 거리가 반경의 2배 안으로 들어오면 15 FPS, 클럭 2400. ④ 경보: 반경 안에 0.5초 머물면 30 FPS로 보면서 알립니다.
경계 구역이 완충 역할을 하기 때문에 감시 상태를 느리게 둬도 안전합니다. 아기가 반경 2배에서 반경 안까지 오는 데 0.3~0.5초가 걸리는데, 그 전에 경계로 올라갑니다.

왼쪽 아래 표가 장면별 실측입니다. 아기가 가만히 있을 때 두 모드는 대기와 고성능을 왔다갔다 하면서 4.2 W였는데 4상태는 1.6 W입니다. 아기가 움직일 때는 4.3에서 2.5 W, 여기에 INT8 모델과 클럭 1500까지 적용하면 1.8 W입니다. 경보 구간은 같은 4상태 FP32가 6.7 W인데 INT8과 클럭 1500을 적용하면 2.6 W입니다. INT8은 안성님이 올린 최신 120에폭 가중치이고, FP32와 같은 물체를 같은 자리에서 찾는 것은 확인했지만 신뢰도가 0.2~0.3 낮아서, 쓰려면 검출 문턱값을 함께 조정해야 합니다.

상태를 몇 개로 할지는 오른쪽 기준으로 정했습니다. 안전은 상태마다 최소 FPS를 정하는 제약이고, 상태 수는 '하나 더 나눠서 얻는 전력이 측정 잡음 0.1 W보다 큰가'로 정합니다. 3에서 4상태로 나눌 때는 칼이 멀리 있는 장면에서 1.1 W가 줄었지만, 4에서 5상태로 원거리 2 FPS 상태를 더 두면 0.08 W라 잡음 이하여서 4에서 멈췄습니다. 하드웨어가 바뀌면 이 경계도 바뀝니다.

단, 아기 장면은 9월 23일 클립으로 만든 영상 입력이라 근사치이고, 월요일 교육장에서 실물로 다시 재서 숫자를 갱신하겠습니다.

[발표자 참고] 4상태 코드(detect_baby_4state.py)와 측정은 9/27 클로드코드가 지인 보드에서 실행. 시간이 부족하면 이 슬라이드는 건너뛰고 앞 슬라이드 끝에 "상태 설계로 아기가 움직여도 1.8 W"만 한 줄 언급. INT8의 정확도(mAP)는 팀원A님 3-2 담당이라 그 숫자를 따를 것. 우리가 확인한 것은 "FP32와 같은 물체를 찾고 신뢰도가 0.2~0.3 낮다"까지.""")

# ═══════════════════ 슬라이드 4a: 결과 분석 ═══════════════════
s = prs.slides.add_slide(BLANK)
header(s, "RESULT ANALYSIS  —  정보윤", "결과 분석 — 목표 대비 달성 여부",
       "목표 4개 중 2개 달성 · 2개 부분 달성(수치 미측정) · 부족한 점 3개는 원인과 개선책을 함께 적었다")
table(s, [["항목", "목표", "결과 (측정 조건)", "판정"],
          ["위험 판정", "아기가 위험물 반경 안에\n0.5초 머물면 경보",
           "4클래스 YOLO11n mAP50-95 0.82 (320 입력, val=test) · 시연 거리 칼 검출률 640 입력 100 % / 320 입력 0 % (9/23 교육장, TS-22)", ("○ 달성\n(640 기준)", C["green"], True)],
          ["오경보 억제", "어른·순간 깜빡임에\n울리지 않음",
           "체류 60 % 판정 + 어른 제외 동작 확인 · 시간당 헛알림 수는 미측정 (집 환경에서 책 → knife 오검출 관찰)", ("△ 부분", C["amber"], True)],
          ["저전력", "평상시 추론 최소화",
           "게이팅 7.12 → 2.52 W (512 FP32, 빈 방, 실제 카메라) · 4상태 평상시 1.6~1.8 W, 경보 구간 6.7 → 2.6 W (320 INT8·클럭 1500, 영상 근사) · PMIC, 웹캠 제외", ("○ 달성", C["green"], True)],
          ["알림", "휴대폰 알림 +\n실시간 화면",
           "ntfy 사진·클립·터널 웹 화면 동작 확인 · 경보 → 휴대폰 도착 지연은 미측정 (9/28 실측)", ("△ 부분", C["amber"], True)]],
      0.23, 1.25, [1.0, 1.85, 5.75, 0.94], size=8.6, hdr_size=8.8, aligns=["l", "l", "l", "c"], heights=[0.27, 0.5, 0.46, 0.56, 0.42])
panel(s, 0.23, 3.6, 9.54, 1.55, C["panel"], adj=0.05)
text(s, 0.38, 3.65, 9.2, 0.24, [[("부족한 점 → 원인 → 개선책", 10, True, C["navy"])]], spacing=1.0)
fix = [("① 데이터", "아기 + 위험물이 함께 있는 학습 사진이 99장뿐 → 동시 장면·배경(칼 없는 방) 사진 보강, val/test 분리"),
       ("② 야간", "일반 웹캠은 빛이 없으면 검출 불가, 학습 사진에도 야간이 없음 → 적외선 LED 카메라 + 야간 사진 학습"),
       ("③ 전체 전력", "PMIC 값에는 USB 웹캠 몫이 없음 → USB-C 전력계로 어댑터 기준 재측정, 장면 비중을 정해 하루 환산")]
for k, (t1, t2) in enumerate(fix):
    l = 0.38 + k * 3.1
    panel(s, l, 3.93, 2.98, 1.12, C["white"], adj=0.1)
    text(s, l + 0.1, 3.98, 2.8, 0.24, [[(t1, 9.5, True, C["blue"])]], spacing=1.0)
    text(s, l + 0.1, 4.24, 2.8, 0.78, [[(t2, 8.6, False, C["gray"])]], spacing=1.05)
note(s, "mAP 는 Colab 시험 사진 227장 · 검출률은 9/23 교육장 실물(칼, 1.5~2 m) · 전력은 라즈베리파이 5 PMIC(USB 웹캠 제외), 지인 보드 9/26~27 · 아기 장면은 클립으로 만든 영상 입력(근사) → 9/28 교육장 실물 재측정")
SCRIPTS[4] = ("4-1. 결과 분석", """4장 결과 분석입니다.
결론은, 목표 네 가지 중 위험 판정과 저전력은 달성했고, 오경보 억제와 알림은 동작은 확인했지만 수치가 빠져서 부분 달성입니다.

표를 보시면, 위험 판정은 4클래스 YOLO11n이 mAP50-95 0.82이고, 시연 거리에서 칼 검출이 640 입력일 때 100 %입니다. 단 320 입력은 먼 칼을 못 잡아서 0 %였고, 검증 데이터와 시험 데이터가 같아서 mAP는 부풀려졌을 수 있습니다.
오경보 억제는 체류 판정과 어른 제외가 동작하는 것은 확인했지만, 시간당 헛알림 횟수는 못 쟀습니다. 집에서는 책을 칼로 잡는 오검출도 봤습니다.
저전력은 게이팅으로 7.12에서 2.52 W, 4상태 설계로 평상시 1.6~1.8 W이고, 전부 보드 전원 칩 PMIC 값이라 웹캠 전력은 빠져 있습니다.
알림은 사진·클립·실시간 화면까지 동작하지만, 경보에서 휴대폰 도착까지의 지연은 실측값을 넣습니다.

부족한 점 세 가지는 원인과 개선책을 같이 적었습니다. 첫째, 아기와 위험물이 함께 있는 학습 사진이 99장뿐이라 데이터 보강이 필요합니다. 둘째, 일반 웹캠이라 밤에는 못 보므로 적외선 카메라와 야간 사진 학습이 필요합니다. 셋째, 웹캠을 포함한 전체 전력은 USB-C 전력계로 다시 재야 합니다.

[발표자 참고] 개선책 세 가지는 팀원A님 5장 향후 과제와 겹칠 수 있으니 발표 전에 안성님 슬라이드와 맞출 것. 9/28 측정값(경보 지연, 헛알림, 실물 전력)이 나오면 표의 "미측정"을 숫자로 교체.""")

# ═══════════════════ 슬라이드 4b: 기대 효과 ═══════════════════
s = prs.slides.add_slide(BLANK)
header(s, "EXPECTED EFFECT  —  정보윤", "기대 효과 — 어디에, 어떻게 쓰나",
       "가정·어린이집에 바로 쓰고, 클래스 설정 2개만 바꾸면 다른 대상으로 확장 · 하루 전력 171 → 60 → 42 Wh (빈 방 기준)")
item(s, 0.23, 1.25, 5.45, 1.12, "01", "가정 — 보호자가 다른 방에 있어도",
     "휴대폰 알림 → 탭하면 실시간 화면 → 앞뒤 5초 클립으로 접근 과정 확인. 인터넷이 끊겨도 보드 LED는 동작. 평상시 영상은 밖으로 나가지 않음")
item(s, 0.23, 2.47, 5.45, 1.12, "02", "어린이집 · 놀이방 — 방마다 보드 1대",
     "원장실 PC 한 화면에 방 4개를 같이 봄 (팀 대시보드 --team 기능 구현됨). 같은 네트워크 안에서 동작, 밖에서 보려면 보드마다 터널 주소")
item(s, 0.23, 3.69, 5.45, 1.12, "03", "대상 · 장소 확장 — 재학습으로",
     "판정 코드는 '감지 주체'와 '위험물' 클래스 설정 2개만 바꾸면 그대로: 노인 낙상 위험 구역, 작업장 위험 설비, 반려동물과 화구. 대상마다 사진 수백 장과 재학습 필요")
text(s, 5.85, 1.25, 3.9, 0.22, [[("하루 전력 환산 — 빈 방, 실제 카메라, 24시간 기준", 9.2, True, C["navy"])]], spacing=1.0)
hero(s, 5.85, 1.5, 1.25, 0.92, "171 Wh", "게이팅 끔\n7.12 W", nsize=15)
hero(s, 7.18, 1.5, 1.25, 0.92, "60 Wh", "게이팅 켬\n2.52 W", nsize=15)
hero(s, 8.51, 1.5, 1.24, 0.92, "42 Wh", "4상태 + 클럭 1500\n대기 1.74 W", nsize=15)
text(s, 5.85, 2.55, 3.9, 0.22, [[("시중 제품과의 차이 — 가장 가까운 CuboAi Danger Zone", 9.2, True, C["navy"])]], spacing=1.0)
table(s, [["", "CuboAi Danger Zone", "우리 시스템"],
          ["위험 구역", "사람이 앱에서 그림", ("위험물을 스스로 찾음", C["blue"], True)],
          ["누가 들어오면", "누구든 알림 (어른 포함)", ("아기만 판정, 어른 제외", C["blue"], True)],
          ["영상", "앱·클라우드 서비스 필요", ("보드 밖으로 안 나감", C["blue"], True)],
          ["구독", "일부 기능 유료", "없음"],
          ["가격", "약 289 달러", "부품 약 17~27만 원 (시제품)"]],
      5.85, 2.8, [0.9, 1.5, 1.5], size=8, hdr_size=8.2, aligns=["l", "l", "l"], heights=[0.26, 0.34, 0.34, 0.34, 0.3, 0.34])
note(s, "하루 환산 = 장면 전력 × 24 h 단순 곱, 아기 활동 시간 비중은 가정하지 않음 (PMIC, 웹캠 제외) · 시중 제품 소비전력은 미공개(어댑터 정격 5 V 2 A = 최대 10 W 만 공개)라 전력 우위는 주장하지 않음 · "
        "제품 정보는 2026-09-27 CuboAi 공식 도움말 기준")
SCRIPTS[5] = ("4-2. 기대 효과", """기대 효과입니다.
결론은, 가정과 어린이집에는 지금 형태로 바로 쓸 수 있고, 클래스 설정 두 개만 바꾸면 다른 대상으로 넓힐 수 있다는 것입니다.

첫째, 가정에서 보호자가 다른 방에 있어도 휴대폰 알림을 받고, 탭하면 실시간 화면이 열리고, 앞뒤 5초 클립으로 아기가 어떻게 접근했는지 확인할 수 있습니다. 평상시 영상은 밖으로 나가지 않습니다.
둘째, 어린이집이나 놀이방은 방마다 보드 한 대를 두고 원장실 PC 한 화면에서 방 네 개를 같이 봅니다. 팀 대시보드 기능이 이미 구현돼 있습니다. 같은 네트워크 안에서 동작한다는 조건은 있습니다.
셋째, 판정 코드는 감지 주체와 위험물 클래스 설정 두 개만 바꾸면 그대로 쓸 수 있어서, 노인 낙상 위험 구역, 작업장 위험 설비, 반려동물과 화구 같은 곳으로 넓힐 수 있습니다. 대상마다 사진 수백 장과 재학습이 필요합니다.

오른쪽 위는 하루 전력 환산입니다. 빈 방 기준으로 게이팅 전 171 Wh가 게이팅으로 60 Wh, 4상태와 클럭 1500으로 42 Wh가 됩니다. 아기가 활동하는 시간 비중은 가정하지 않은 단순 환산입니다.
오른쪽 아래는 시중 제품과의 차이입니다. 가장 가까운 CuboAi의 Danger Zone은 사람이 앱에서 구역을 그려야 하고 어른이 지나가도 울리는데, 저희는 위험물을 스스로 찾고 아기만 판정하며 영상이 밖으로 나가지 않습니다. 시중 제품은 소비전력을 공개하지 않아서 저희가 더 저전력이라고는 말하지 않겠습니다. 이상입니다.

[발표자 참고] "아기 활동 시간 비중은?"이 나오면 "장면 비중을 정해야 하고, 가정하지 않아서 빈 방 기준으로만 환산했다"고 답. 제품 가격·기능은 2026-09-27 웹 검색(CuboAi 도움말) 기준이고 전수 조사가 아니므로 "저희가 조사한 범위에서는"을 붙일 것.""")

for k, sl in enumerate(prs.slides, 1):
    notes(sl, "[%s]\n\n%s" % (SCRIPTS[k][0], SCRIPTS[k][1]))
final_out = OUT
for _try in range(10):
    try:
        prs.save(final_out); break
    except PermissionError:
        m = re.search(r"_v(\d+)\.pptx$", final_out)
        final_out = final_out[:m.start()] + "_v%d.pptx" % (int(m.group(1)) + 1)
print("saved", final_out, "slides:", len(prs.slides))
with open(os.path.join(FIG, "발표대본_정보윤.md"), "w", encoding="utf-8") as f:
    f.write("# 정보윤 파트 발표 대본 (슬라이드 노트와 동일, 2026-09-27)\n\n")
    f.write("슬라이드 파일: `" + os.path.basename(final_out) + "` (같은 폴더). PowerPoint 에서 보기 → 슬라이드 노트로도 볼 수 있다.\n\n")
    for k in sorted(SCRIPTS):
        f.write("## 슬라이드 %d — %s\n\n%s\n\n" % (k, SCRIPTS[k][0], SCRIPTS[k][1]))
print("script md written")
