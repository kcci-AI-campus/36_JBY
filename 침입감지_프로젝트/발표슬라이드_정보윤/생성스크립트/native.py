# -*- coding: utf-8 -*-
# 슬라이드 안의 흐름도·상태도·그래프를 이미지가 아니라 PowerPoint 도형·연결선·네이티브 차트로 그린다.
# 모든 라벨은 개별 텍스트 상자/도형이라 PowerPoint 에서 끌어 옮길 수 있다. 화살표는 가능한 한 도형에 붙인 연결선.
import json, os
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import XyChartData, CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

G = {}
BLUE, NAVY, GRAY, MUTED, LGRAY = "2563EB", "0F172A", "475569", "64748B", "94A3B8"
PANEL, WHITE, RED, GREEN, AMBER, ORANGE = "F1F5F9", "FFFFFF", "DC2626", "16A34A", "D97706", "F97316"


def init(g):
    G.update(g)


def _run(p, txt, size, bold, color):
    r = p.add_run(); r.text = txt
    f = r.font; f.name = G["FONT"]; f.size = Pt(size); f.bold = bold; f.color.rgb = G["rgb"](color)
    G["set_ea"](r)


def sbox(s, l, t, w, h, title, sub="", ec=BLUE, fill=WHITE, tc=NAVY, ts=9.8, ss=7.9, lw=1.75, adj=0.12,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, wrap=True, name=None):
    sh = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = adj
    sh.fill.solid(); sh.fill.fore_color.rgb = G["rgb"](fill)
    if ec:
        sh.line.color.rgb = G["rgb"](ec); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if name: sh.name = name
    tf = sh.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.0
    _run(p, title, ts, True, tc)
    for line in (sub.split("\n") if sub else []):
        q = tf.add_paragraph(); q.alignment = PP_ALIGN.CENTER; q.line_spacing = 1.0
        _run(q, line, ss, False, GRAY)
    return sh


def dot(s, cx, cy, label, color, r=0.11, size=7.5):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    sh.fill.solid(); sh.fill.fore_color.rgb = G["rgb"](color)
    sh.line.color.rgb = G["rgb"](WHITE); sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, label, size, True, WHITE)
    return sh


def _arrowhead(conn, color, lw):
    conn.line.color.rgb = G["rgb"](color); conn.line.width = Pt(lw)
    ln = conn.line._get_or_add_ln()
    for old in ln.findall(qn("a:tailEnd")): ln.remove(old)
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


def link(s, a, ai, b, bi, color=LGRAY, lw=1.5):
    """도형 a 의 연결점 ai → 도형 b 의 연결점 bi (0 위, 1 왼쪽, 2 아래, 3 오른쪽). 도형을 옮기면 화살표가 따라간다."""
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(a, ai); c.end_connect(b, bi)
    _arrowhead(c, color, lw)
    return c


def arrow(s, x1, y1, x2, y2, color=LGRAY, lw=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _arrowhead(c, color, lw)
    return c


def label(s, l, t, w, h, txt, size=7.5, bold=True, color=GRAY, align="c"):
    return G["text"](s, l, t, w, h, [[(txt, size, bold, color)]], align=align, anchor="m", spacing=1.0, margin=0.0)


# ───────── 슬라이드 1: 알림 파이프라인 ─────────
def pipeline(s):
    label(s, 0.35, 1.29, 9.3, 0.2, "한 프레임의 흐름 — 전부 보드 안에서 처리. 평상시 영상은 밖으로 나가지 않고, 경보 순간의 사진·클립만 나간다",
          size=9.2, color=NAVY, align="l")
    steps = [("① 카메라 → YOLO11n 추론", "USB 웹캠 640×480 · 30 FPS\nbaby·adult·knife·outlet 검출\nLiteRT 로 보드에서 추론", BLUE, NAVY),
             ("② 거리 판정", "위험 반경 = 위험물 박스 폭 × 1.5\n아기 박스가 반경에 닿으면 '위험'", BLUE, NAVY),
             ("③ 체류 판정", "최근 0.5초 프레임의 60 % 이상 위험\n(최소 3장) 이면 경보\n깜빡이는 오검출 제거", BLUE, NAVY),
             ("④ 경보 출력", "보드 LED 즉시 빨강\n휴대폰 알림음 · 웹으로 전달 (아래)", RED, RED)]
    top = []
    for k, (t_, s_, ec, tc) in enumerate(steps):
        top.append(sbox(s, 0.35 + 2.4 * k, 1.55, 2.2, 0.98, t_, s_, ec=ec, tc=tc, name="단계%d" % (k + 1)))
    for k in range(3):
        link(s, top[k], 3, top[k + 1], 1, color=BLUE, lw=1.75)
    outs = [("보드 LED", "초록·파랑·빨강 3색, 즉시", GREEN),
            ("휴대폰 알림 (ntfy)", "알림음 + 사진 1장 + 앞 3초·뒤 2초 클립\n같은 경보는 최소 30초 간격", AMBER),
            ("실시간 화면", "알림을 탭하면 웹 페이지\nCloudflare 터널 + 접속 암호", BLUE),
            ("웹 대시보드", "상태·경보 기록·LED·FPS·전력\n(3-4 팀원B 파트)", GRAY)]
    for k, (t_, s_, c) in enumerate(outs):
        o = sbox(s, 0.35 + 2.4 * k, 3.1, 2.2, 0.95, t_, s_, ec=c, fill=PANEL, tc=c, ts=9.6, lw=1.5, name="출력%d" % (k + 1))
        link(s, top[3], 2, o, 0, color=RED, lw=1.25)
    label(s, 8.75, 2.68, 0.95, 0.18, "경보 발생", size=8.2, color=RED, align="r")


# ───────── 슬라이드 2: 게이팅 흐름도 ─────────
def gating(s):
    f = sbox(s, 1.13, 1.375, 1.55, 0.36, "새 프레임 (카메라)", ec=LGRAY, ts=8.8, lw=1.5)
    a = sbox(s, 0.83, 1.955, 2.15, 0.58, "직전 프레임과 차이 계산", "160×120 흑백으로 줄여\n'변한 픽셀 비율'만 계산",
             ec=BLUE, ts=8.8, ss=6.9, lw=1.5)
    d = sbox(s, 1.1, 2.735, 1.6, 0.94, "변한 픽셀 < 0.6 % ?", "그리고 위험 중 아님", ec=AMBER, ts=8.2, ss=6.9, lw=1.5,
             shape=MSO_SHAPE.DIAMOND, wrap=False)
    inf = sbox(s, 0.3, 2.855, 0.55, 0.7, "추론", "YOLO11n\n실행", ec=RED, fill="FEF2F2", tc=RED, ts=7.6, ss=6.2, lw=1.25)
    skip = sbox(s, 2.93, 2.855, 0.55, 0.7, "건너뜀", "직전 결과\n재사용", ec=GREEN, fill="ECFDF5", tc=GREEN, ts=7.6, ss=6.2, lw=1.25)
    link(s, f, 2, a, 0); link(s, a, 2, d, 0)
    link(s, d, 1, inf, 3, color=RED); link(s, d, 3, skip, 1, color=GREEN)
    label(s, 0.82, 3.0, 0.32, 0.16, "아니오", size=7, color=RED)
    label(s, 2.7, 3.0, 0.22, 0.16, "예", size=7, color=GREEN)
    dot(s, 0.85, 1.955, "A", BLUE); dot(s, 1.2, 2.83, "B", AMBER)
    G["panel"](s, 0.33, 3.725, 3.15, 0.6, PANEL, adj=0.1)
    dot(s, 0.48, 4.025, "C", GRAY)
    G["text"](s, 0.64, 3.75, 2.8, 0.55, [[("안전장치  ① 위험 중이면 절대 건너뛰지 않음", 6.9, False, GRAY)],
                                        [("② 연속 15장 건너뛰면 강제로 한 번 추론", 6.9, False, GRAY)],
                                        [("③ 쉴 때도 8 FPS 로 확인 (최대 공백 0.13초)", 6.9, False, GRAY)]], anchor="m", spacing=1.05)


# ───────── 슬라이드 2: 게이팅 타임라인 (측정 trace → 네이티브 산점도 차트) ─────────
def _load(path):
    ev = json.load(open(path, encoding="utf-8"))["traceEvents"]
    t0 = min(e["ts"] for e in ev if e.get("ph") in ("X", "C"))
    inf = [(e["ts"] - t0) / 1e6 for e in ev if e.get("ph") == "X" and e["name"] == "inference"]
    cpu = [((e["ts"] - t0) / 1e6, e["args"]["total"]) for e in ev if e.get("ph") == "C" and e["name"] == "CPU total %"]
    return inf, cpu


def _xy_chart(s, l, t, w, h, inf, cpu, show_x):
    cd = XyChartData()
    s1 = cd.add_series("CPU 사용률 %")
    for x, y in cpu:
        if x <= 60: s1.add_data_point(round(x, 3), round(y, 1))
    s2 = cd.add_series("추론 실행 순간")
    for x in inf:
        if x <= 60: s2.add_data_point(round(x, 3), 12)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(l), Inches(t), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.size = Pt(7); ch.font.name = G["FONT"]; ch.font.color.rgb = G["rgb"](MUTED)
    ser_cpu, ser_inf = ch.plots[0].series[0], ch.plots[0].series[1]
    ser_cpu.format.line.color.rgb = G["rgb"](BLUE); ser_cpu.format.line.width = Pt(0.9); ser_cpu.smooth = False
    ser_inf.format.line.fill.background(); ser_inf.smooth = False
    ser_inf.marker.style = XL_MARKER_STYLE.SQUARE; ser_inf.marker.size = 3
    ser_inf.marker.format.fill.solid(); ser_inf.marker.format.fill.fore_color.rgb = G["rgb"](ORANGE)
    ser_inf.marker.format.line.fill.background()
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, 60, 10
    ya.minimum_scale, ya.maximum_scale, ya.major_unit = 0, 100, 50
    ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = G["rgb"]("E2E8F0")
    xa.has_major_gridlines = False
    for ax in (xa, ya):
        ax.format.line.color.rgb = G["rgb"]("E2E8F0")
        ax.tick_labels.font.size = Pt(7); ax.tick_labels.font.color.rgb = G["rgb"](MUTED)
    if not show_x:
        xa.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    return gf


def timeline(s, base):
    off_inf, off_cpu = _load(os.path.join(base, "trace_gate_off.json"))
    on_inf, on_cpu = _load(os.path.join(base, "trace_gate_on.json"))
    label(s, 3.72, 1.27, 6.0, 0.17, "같은 장면·같은 모델(yolo11n 512)·60초, 라즈베리파이 5, 2026-09-26 — 추론 실행 순간과 CPU 사용률",
          size=7.6, bold=False, color=MUTED, align="l")
    label(s, 3.72, 1.46, 6.0, 0.19, "Before — 게이팅 끔:  7.12 W · CPU 92 % · 60초 동안 추론 749회",
          size=8.6, color=NAVY, align="l")
    _xy_chart(s, 3.68, 1.63, 6.09, 1.12, off_inf, off_cpu, show_x=False)
    label(s, 3.72, 2.78, 6.0, 0.19, "After — 게이팅 켬:  2.52 W · CPU 5 % · 60초 동안 추론 32회 (93 % 건너뜀)",
          size=8.6, color=NAVY, align="l")
    _xy_chart(s, 3.68, 2.95, 6.09, 1.25, on_inf, on_cpu, show_x=True)
    label(s, 3.72, 4.2, 6.0, 0.17, "주황 점 = 추론이 실제로 실행된 순간 · 파란 선 = CPU 전체 사용률(4코어 평균, %) · 가로축 = 시간(초)",
          size=7.2, bold=False, color=MUTED, align="l")


# ───────── 슬라이드 3: 4상태 상태도 ─────────
def states(s):
    label(s, 0.33, 1.32, 6.25, 0.2, "상태마다 '얼마나 자주 · 어떤 클럭으로' 추론할지를 다르게 — 경계 구역(반경×2)이 완충 역할",
          size=8.6, bold=False, color=GRAY)
    st = [("① 대기", "아기 없음\n5초에 1장\n클럭 1500", BLUE),
          ("② 감시", "아기 있음, 위험물 멀리\n5 FPS · 게이팅\n클럭 1500", GREEN),
          ("③ 경계", "거리 < 반경 × 2\n15 FPS\n클럭 2400", AMBER),
          ("④ 경보", "반경 안 0.5초 체류\n30 FPS · 클럭 2400\nLED · 휴대폰 알림", RED)]
    ups = ["아기 보임", "위험물에 접근", "반경 안 0.5초"]
    downs = ["아기 3초 안 보임", "멀어진 지 3초", "위험 없어진 지 3초"]
    for k, (t_, s_, c) in enumerate(st):
        sbox(s, 0.43 + 1.57 * k, 1.855, 1.32, 0.95, t_, s_, ec=c, tc=c, ts=10.5, ss=7.0, lw=1.9, adj=0.1)
    for k in range(3):
        x1 = 0.43 + 1.57 * k + 1.32; x2 = x1 + 0.25
        c = st[k + 1][2]
        arrow(s, x1, 2.14, x2, 2.14, color=c, lw=1.6)
        arrow(s, x2, 2.52, x1, 2.52, color=LGRAY, lw=1.3)
        label(s, x1 - 0.35, 1.62, 0.95, 0.17, ups[k], size=7.2, color=c)
        label(s, x1 - 0.4, 2.84, 1.05, 0.17, downs[k], size=6.8, bold=False, color=MUTED)
    label(s, 0.33, 3.1, 6.25, 0.22, "원리: 전력 = 유휴 + 추론 1회 에너지 × 초당 횟수  →  위험이 멀수록 천천히·싸게, 가까울수록 빠르게·정확하게",
          size=8.4, color=NAVY)


# ───────── 슬라이드 3: 두 장면 × 세 설계 막대 (네이티브 차트) ─────────
def bars(s):
    cd = CategoryChartData()
    cd.categories = ["아기 움직임, 위험물 없음", "아기+칼 반경 안 (경보 구간)"]
    series = [("두 모드", (4.29, 7.17), LGRAY), ("4상태", (2.53, 6.66), BLUE), ("4상태+INT8+클럭 1500", (1.78, 2.58), ORANGE)]
    for n, v, _ in series: cd.add_series(n, v)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.5), Inches(3.58), Inches(3.27), Inches(1.6), cd)
    ch = gf.chart
    ch.font.size = Pt(7); ch.font.name = G["FONT"]; ch.font.color.rgb = G["rgb"](GRAY)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(6.5)
    pl = ch.plots[0]; pl.gap_width = 60; pl.overlap = 0
    pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = "0.0"; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(7); dl.font.bold = True; dl.font.color.rgb = G["rgb"](NAVY)
    for ser, (_, _, c) in zip(pl.series, series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = G["rgb"](c); ser.format.line.fill.background()
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = 0, 8.5
    va.has_major_gridlines = False; va.visible = False
    ca = ch.category_axis; ca.format.line.color.rgb = G["rgb"]("E2E8F0")
    ca.tick_labels.font.size = Pt(7); ca.tick_labels.font.color.rgb = G["rgb"](NAVY)
    label(s, 6.55, 3.6, 1.5, 0.15, "평균 전력 (W)", size=6.8, bold=False, color=MUTED, align="l")
