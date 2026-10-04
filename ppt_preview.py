# -*- coding: utf-8 -*-
"""만든 PPT의 내용을 HTML 한 장으로 뽑아 본다 (폰에서 확인용)."""
import os, html
from pptx import Presentation
from pptx.util import Emu

SRC = 'AI_시스템반도체_설계_2기_결과보고서(정보윤).pptx'
OUT = 'ppt_미리보기.html'

prs = Presentation(SRC)
parts = []

for i, slide in enumerate(prs.slides, 1):
    texts, tables, pics = [], [], []
    for sh in slide.shapes:
        if sh.shape_type == 13 or sh.__class__.__name__ == 'Picture':
            pics.append(sh)
            continue
        if sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append([c.text_frame.text.strip() for c in r.cells])
            tables.append(rows)
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            for p in sh.text_frame.paragraphs:
                t = ''.join(r.text for r in p.runs).strip()
                if t:
                    big = any(r.font.size and r.font.size.pt >= 20 for r in p.runs)
                    bold = any(r.font.bold for r in p.runs)
                    texts.append((t, big, bold))

    body = []
    for t, big, bold in texts:
        cls = 'big' if big else ('b' if bold else '')
        body.append('<p class="%s">%s</p>' % (cls, html.escape(t)))
    for rows in tables:
        h = ['<table>']
        for j, row in enumerate(rows):
            tag = 'th' if j == 0 else 'td'
            h.append('<tr>' + ''.join('<%s>%s</%s>' % (tag, html.escape(c), tag)
                                      for c in row) + '</tr>')
        h.append('</table>')
        body.append(''.join(h))
    if pics:
        body.append('<p class="pic">[그림 %d장 포함]</p>' % len(pics))

    parts.append('<section><h2>슬라이드 %d</h2>%s</section>' % (i, ''.join(body)))

CSS = """
:root{--bg:#fff;--fg:#1a1a1a;--muted:#666;--line:#d8dee6;--navy:#1f3b63;--accent:#2e6db4;--soft:#eef3f9}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#15181c;--fg:#e8eaed;--muted:#a0a6ad;--line:#333a42;--navy:#7fa8dd;--accent:#8fb6e8;--soft:#1e242c}}
:root[data-theme="dark"]{--bg:#15181c;--fg:#e8eaed;--muted:#a0a6ad;--line:#333a42;--navy:#7fa8dd;--accent:#8fb6e8;--soft:#1e242c}
body{background:var(--bg);color:var(--fg);font-family:'Malgun Gothic','맑은 고딕',system-ui,sans-serif;
     line-height:1.6;margin:0;padding:16px;max-width:900px;margin:0 auto}
h1{font-size:1.3rem;color:var(--navy);border-bottom:3px solid var(--accent);padding-bottom:8px}
section{border:1px solid var(--line);border-radius:8px;padding:14px 16px;margin:18px 0;background:var(--bg)}
h2{font-size:.85rem;color:var(--muted);margin:0 0 10px;letter-spacing:.04em;text-transform:uppercase}
p{margin:6px 0;font-size:.92rem}
p.big{font-size:1.15rem;font-weight:700;color:var(--navy);margin:10px 0}
p.b{font-weight:700;color:var(--accent)}
p.pic{color:var(--muted);font-style:italic;font-size:.85rem}
.tblwrap,table{width:100%}
table{border-collapse:collapse;margin:10px 0;font-size:.82rem;display:block;overflow-x:auto;white-space:nowrap}
th{background:var(--navy);color:#fff;padding:6px 9px;text-align:left;font-weight:700}
td{border-bottom:1px solid var(--line);padding:6px 9px}
tr:nth-child(even) td{background:var(--soft)}
"""

io_out = ('<title>결과보고서 미리보기</title><style>%s</style>'
          '<h1>AI 시스템반도체 설계 2기 · 결과보고서 (슬라이드 %d장)</h1>%s'
          % (CSS, len(prs.slides._sldIdLst), ''.join(parts)))
open(OUT, 'w', encoding='utf-8').write(io_out)
print('생성:', os.path.abspath(OUT))
